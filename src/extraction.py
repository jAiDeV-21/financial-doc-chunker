import logging
import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from config import ChunkConfig
from data_models import Block, ParsedDoc
from text_utilities import (
    _HEAD_NUM,
    _LIST_MARK,
    _PAGE_NO,
    clean_text,
    join_lines,
    looks_like_heading,
    numbering_depth,
)

log = logging.getLogger(__name__)


@dataclass
class _Seg:
    text: str
    size: float
    bold: bool
    bbox: tuple
    page: int
    kind: str = "text"  # text | table
    table: Optional[tuple] = None
    cls: str = "paragraph"  # heading | list_item | paragraph | footnote | title
    sig: tuple = ()
    level: int = 0
    toc_level: int = 0
    plain: bool = False  # no font info (OCR) -> rely on text cues only


def _is_num(c: str) -> bool:
    if re.fullmatch(r"(?:19|20)\d{2}", c):
        return False  # a year is a header, not data
    return bool(re.fullmatch(r"[(\-–]?[$€£₹]?\s?[\d,]+(?:\.\d+)?\s?%?\)?", c))


def _clean_rows(raw) -> list:
    rows = [[clean_text((c or "").replace("\n", " ")) for c in r] for r in raw]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    keep = [j for j in range(ncol) if any(r[j] for r in rows)]
    return [[r[j] for j in keep] for r in rows]


def split_table_header(rows: list) -> tuple[list, list]:
    first = [c for c in rows[0] if c]
    if first and sum(_is_num(c) for c in first) / len(first) < 0.5:
        return rows[0], rows[1:]
    return [""] * len(rows[0]), rows


def _find_tables(page, strategy: str) -> list[_Seg]:
    out = []
    try:
        finder = page.find_tables(strategy=strategy)
    except Exception as e:  # noqa: BLE001
        log.debug("find_tables failed on page %s: %s", page.number + 1, e)
        return out
    for t in finder.tables:
        rows = _clean_rows(t.extract())
        if len(rows) < 2 or len(rows[0]) < 2:
            continue
        out.append(
            _Seg(
                "",
                0,
                False,
                tuple(t.bbox),
                page.number + 1,
                kind="table",
                table=split_table_header(rows),
            )
        )
    return out


def _inside(b: tuple, boxes: list, frac: float = 0.6) -> bool:
    area = max((b[2] - b[0]) * (b[3] - b[1]), 1e-6)
    for t in boxes:
        w = min(b[2], t[2]) - max(b[0], t[0])
        h = min(b[3], t[3]) - max(b[1], t[1])
        if w > 0 and h > 0 and (w * h) / area >= frac:
            return True
    return False


def _block_to_segs(blk: dict, pno: int) -> list[_Seg]:
    """One PyMuPDF text block -> segments of uniform style; new segment at each list marker."""
    groups: list[dict] = []
    for ln in blk.get("lines", []):
        spans = [s for s in ln["spans"] if s["text"].strip()]
        if not spans:
            continue
        text = "".join(s["text"] for s in ln["spans"])
        n = sum(len(s["text"].strip()) for s in spans)
        size_chars: Counter = Counter()
        bold_n = 0
        for s in spans:
            size_chars[round(s["size"] * 2) / 2] += len(s["text"].strip())
            if (s["flags"] & 16) or re.search(r"bold|black|heavy", s["font"], re.I):
                bold_n += len(s["text"].strip())
        size = size_chars.most_common(1)[0][0]
        bold = bold_n > 0.5 * n
        g = groups[-1] if groups else None
        if (
            g is None
            or abs(g["size"] - size) > 0.6
            or g["bold"] != bold
            or _LIST_MARK.match(text)
        ):
            g = {"size": size, "bold": bold, "lines": [], "bbox": list(ln["bbox"])}
            groups.append(g)
        g["lines"].append(text)
        g["bbox"] = [
            min(g["bbox"][0], ln["bbox"][0]),
            min(g["bbox"][1], ln["bbox"][1]),
            max(g["bbox"][2], ln["bbox"][2]),
            max(g["bbox"][3], ln["bbox"][3]),
        ]
    segs = []
    for g in groups:
        t = join_lines(g["lines"])
        if t:
            segs.append(_Seg(t, g["size"], g["bold"], tuple(g["bbox"]), pno))
    return segs


def _reading_order(items: list[_Seg], page_w: float) -> list[_Seg]:
    """Top-to-bottom, but two-column pages are read left column then right column."""

    def y(i):
        return (round(i.bbox[1]), i.bbox[0])

    mid, tol = page_w / 2, 0.03 * page_w
    full = sorted((i for i in items if i.bbox[2] - i.bbox[0] > 0.6 * page_w), key=y)
    narrow = [i for i in items if i not in full]
    left = [i for i in narrow if i.bbox[2] <= mid + tol]
    right = [i for i in narrow if i.bbox[0] >= mid - tol]
    if not (
        len(left) >= 3 and len(right) >= 3 and len(left) + len(right) == len(narrow)
    ):
        return sorted(items, key=y)

    def band(i):
        return sum(1 for f in full if f.bbox[3] <= i.bbox[1] + 2)

    out: list[_Seg] = []
    for b in range(len(full) + 1):
        out += sorted((i for i in left if band(i) == b), key=y)
        out += sorted((i for i in right if band(i) == b), key=y)
        if b < len(full):
            out.append(full[b])
    return out


def _ocr_page(page, pno: int, size: float) -> list[_Seg]:
    import io

    import pytesseract
    from PIL import Image

    img = Image.open(io.BytesIO(page.get_pixmap(dpi=200).tobytes("png")))
    segs, y = [], 0.0
    for para in re.split(r"\n\s*\n", pytesseract.image_to_string(img)):
        t = join_lines(para.splitlines())
        if t:
            segs.append(
                _Seg(t, size, False, (0, y, page.rect.width, y + 1), pno, plain=True)
            )
            y += 2
    return segs


def _norm_title(t: str) -> str:
    return re.sub(r"\W+", " ", t.lower()).strip()


def extract_pdf(path: Path, cfg: ChunkConfig) -> ParsedDoc:
    import pymupdf

    doc = pymupdf.open(str(path))
    n_pages = doc.page_count
    pages = []  # (page_w, page_h, [segs])
    ocr_pages = 0
    for pno in range(n_pages):
        page = doc[pno]
        W, H = page.rect.width, page.rect.height
        tables = _find_tables(page, cfg.table_strategy)
        tboxes = [t.bbox for t in tables]
        segs: list[_Seg] = []
        for blk in page.get_text("dict")["blocks"]:
            if blk.get("type") == 0 and not _inside(blk["bbox"], tboxes):
                segs += _block_to_segs(blk, pno + 1)
        if cfg.ocr and sum(len(s.text) for s in segs) < 20 and not tables:
            segs, ocr_pages = _ocr_page(page, pno + 1, 10.0), ocr_pages + 1
        pages.append((W, H, segs, tables))

    # ---- body font size = the size carrying most characters
    size_chars: Counter = Counter()
    for _, _, segs, _ in pages:
        for s in segs:
            size_chars[s.size] += len(s.text)
    body = size_chars.most_common(1)[0][0] if size_chars else 10.0

    # ---- drop running headers/footers and bare page numbers
    if not cfg.keep_headers_footers:
        key = lambda s: re.sub(r"\d+", "#", s.text.lower())  # noqa: E731
        margin_pages: dict[str, set] = {}
        for pi, (W, H, segs, _) in enumerate(pages):
            for s in segs:
                if s.bbox[1] < 0.07 * H or s.bbox[3] > 0.93 * H:
                    margin_pages.setdefault(key(s), set()).add(pi)
        repeated = {
            k for k, v in margin_pages.items() if len(v) >= max(3, 0.3 * n_pages)
        }
        removed = 0
        for pi, (W, H, segs, tb) in enumerate(pages):
            kept = []
            for s in segs:
                in_margin = s.bbox[1] < 0.07 * H or s.bbox[3] > 0.93 * H
                if in_margin and (key(s) in repeated or _PAGE_NO.match(s.text)):
                    removed += 1
                    continue
                kept.append(s)
            pages[pi] = (W, H, kept, tb)
        log.debug("removed %d header/footer segments", removed)

    # ---- classify every segment
    all_segs: list[_Seg] = []
    for W, H, segs, _ in pages:
        for s in segs:
            ratio = s.size / body if body else 1.0
            if (
                s.size < 0.85 * body
                and s.bbox[1] > 0.75 * H
                and not _LIST_MARK.match(s.text)
            ):
                s.cls = "footnote"
            elif looks_like_heading(s.text, s.bold, ratio, plain=s.plain):
                s.cls, s.sig = "heading", (s.size, s.bold)
            elif _LIST_MARK.match(s.text):
                s.cls = "list_item"
            all_segs.append(s)

    # ---- PDF bookmarks (outline), when they match the page text, are ground truth for levels
    toc_used = False
    try:
        toc = doc.get_toc(simple=True)
    except Exception:  # noqa: BLE001
        toc = []
    if len(toc) >= 3:
        toc_map = {}
        for lvl, title, pg in toc:
            toc_map.setdefault(_norm_title(title), (lvl, pg))
        hits = 0
        for s in all_segs:
            hit = toc_map.get(_norm_title(s.text))
            if hit and abs(hit[1] - s.page) <= 1 and len(s.text) < 200:
                s.cls, s.toc_level, s.sig = "heading", hit[0], s.sig or (s.size, s.bold)
                hits += 1
        toc_used = hits >= 0.5 * len(toc)
        if not toc_used:
            for s in all_segs:
                s.toc_level = 0
    heads = [s for s in all_segs if s.cls == "heading"]

    # ---- document title: a unique, top-ranked heading on the first pages
    meta_title = ((doc.metadata or {}).get("title") or "").strip()
    title = (
        path.stem
        if len(meta_title) < 3 or meta_title.lower() in {"(anonymous)", "untitled"}
        else meta_title
    )
    free = [s for s in heads if not s.toc_level]
    if free:
        top = max(s.sig for s in free)
        top_segs = [s for s in free if s.sig == top]
        if len(top_segs) == 1 and top_segs[0].page <= 2:
            top_segs[0].cls = "title"
            title = top_segs[0].text
            heads = [s for s in heads if s is not top_segs[0]]

    # ---- heading levels
    sigs = sorted(
        {s.sig for s in heads if not s.toc_level}, key=lambda g: (-g[0], -g[1])
    )
    rank = {g: i + 1 for i, g in enumerate(sigs)}
    max_toc = max((s.toc_level for s in heads), default=0)
    for s in heads:
        if s.toc_level:
            s.level = s.toc_level
        else:
            s.level = rank[s.sig] + (max_toc if toc_used else 0)
            if len(sigs) == 1 and not toc_used and numbering_depth(s.text):
                s.level = numbering_depth(s.text)

    # ---- assemble blocks in reading order
    blocks: list[Block] = []
    for W, H, segs, tables in pages:
        for s in _reading_order(segs + tables, W):
            if s.kind == "table":
                header, rows = s.table
                blocks.append(
                    Block(
                        "table", page=s.page, page_end=s.page, header=header, rows=rows
                    )
                )
            elif s.cls == "title":
                continue
            elif s.cls == "heading":
                blocks.append(Block("heading", s.text, s.page, s.page, level=s.level))
            elif s.cls == "footnote":
                blocks.append(Block("footnote", s.text, s.page, s.page))
            elif s.cls == "list_item":
                blocks.append(Block("_li", s.text, s.page, s.page))
            else:
                blocks.append(Block("paragraph", s.text, s.page, s.page))

    blocks = _merge_across_pages(blocks)
    blocks = _group_lists(blocks)
    blocks = _attach_captions(blocks)
    info = {
        "extractor": "pymupdf",
        "body_font_size": body,
        "toc_used": toc_used,
        "ocr_pages": ocr_pages,
        "n_tables": sum(b.kind == "table" for b in blocks),
    }
    return ParsedDoc(title, blocks, n_pages, info)


def _merge_across_pages(blocks: list[Block]) -> list[Block]:
    """Stitch paragraphs / tables that continue over a page break."""
    out: list[Block] = []
    for b in blocks:
        j = len(out) - 1
        while j >= 0 and out[j].kind == "footnote":
            j -= 1
        prev = out[j] if j >= 0 else None
        if (
            prev
            and b.kind == "paragraph"
            and prev.kind == "paragraph"
            and b.page == (prev.page_end or 0) + 1
            and b.text[:1].islower()
            and prev.text[-1:] not in '.!?:;"”)'
        ):
            prev.text = join_lines([prev.text, b.text])
            prev.page_end = b.page
            continue
        if (
            prev
            and b.kind == "table"
            and prev.kind == "table"
            and len(prev.header) == len(b.header)
            and b.page == (prev.page_end or 0) + 1
        ):
            rows = (
                b.rows[1:]
                if b.rows
                and b.header == [""] * len(b.header)
                and prev.header != [""] * len(prev.header)
                and b.rows[0] == prev.header
                else b.rows
            )
            if b.header != [""] * len(b.header) and b.header != prev.header:
                rows = [b.header] + rows  # new header text that is really data
            prev.rows += rows
            prev.page_end = b.page
            continue
        out.append(b)
    return out


def _group_lists(blocks: list[Block]) -> list[Block]:
    """Consecutive list items (+ a lead-in paragraph ending in ':') become one list block."""
    out: list[Block] = []
    i = 0
    while i < len(blocks):
        b = blocks[i]
        if b.kind != "_li":
            out.append(b)
            i += 1
            continue
        j = i
        while j < len(blocks) and blocks[j].kind == "_li":
            j += 1
        run = blocks[i:j]
        if len(run) == 1:
            out.append(Block("paragraph", run[0].text, run[0].page, run[0].page_end))
        else:
            intro = ""
            if (
                out
                and out[-1].kind == "paragraph"
                and out[-1].text.rstrip().endswith(":")
            ):
                lead = out.pop()
                intro, page = lead.text, lead.page
            else:
                page = run[0].page
            out.append(
                Block(
                    "list",
                    page=page,
                    page_end=run[-1].page_end,
                    intro=intro,
                    items=[r.text for r in run],
                )
            )
        i = j
    return out


_CAPTION = re.compile(r"^(?:table|exhibit|schedule|figure)\s+[\w.-]+", re.I)


def _attach_captions(blocks: list[Block]) -> list[Block]:
    out: list[Block] = []
    for b in blocks:
        if (
            b.kind == "table"
            and out
            and out[-1].kind == "paragraph"
            and len(out[-1].text) < 200
            and _CAPTION.match(out[-1].text)
        ):
            b.caption = out.pop().text
        out.append(b)
    return out


# ---------------------------------------------------------------- DOCX / TXT / MD
def extract_docx(path: Path, cfg: ChunkConfig) -> ParsedDoc:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = Document(str(path))
    title = (d.core_properties.title or "").strip() or path.stem
    has_styles = any(
        (p.style.name or "").lower().startswith(("heading", "title"))
        for p in d.paragraphs
    )
    blocks: list[Block] = []
    for child in d.element.body.iterchildren():
        if child.tag == qn("w:p"):
            p = Paragraph(child, d)
            t = clean_text(p.text)
            if not t:
                continue
            style = (p.style.name or "") if p.style is not None else ""
            m = re.match(r"heading\s*(\d)", style, re.I)
            runs = [r for r in p.runs if r.text.strip()]
            bold = bool(runs) and all(r.bold for r in runs)
            if style.lower() == "title":
                title = t
            elif m:
                blocks.append(Block("heading", t, level=int(m.group(1))))
            elif not has_styles and looks_like_heading(t, bold):
                blocks.append(Block("heading", t, level=numbering_depth(t) or 2))
            elif "list" in style.lower() or (
                p._p.pPr is not None and p._p.pPr.numPr is not None
            ):
                blocks.append(Block("_li", t if _LIST_MARK.match(t) else f"- {t}"))
            elif _LIST_MARK.match(t) and len(t) < 600:
                blocks.append(Block("_li", t))
            else:
                blocks.append(Block("paragraph", t))
        elif child.tag == qn("w:tbl"):
            tbl = Table(child, d)
            raw = []
            for row in tbl.rows:
                seen, cells = set(), []
                for c in row.cells:  # merged cells repeat; keep one
                    if id(c._tc) not in seen:
                        seen.add(id(c._tc))
                        cells.append(c.text)
                raw.append(cells)
            rows = _clean_rows(raw)
            if len(rows) >= 2 and len(rows[0]) >= 2:
                header, body = split_table_header(rows)
                blocks.append(Block("table", header=header, rows=body))
    blocks = _attach_captions(_group_lists(blocks))
    return ParsedDoc(
        title, blocks, None, {"extractor": "python-docx", "styled_headings": has_styles}
    )


def extract_text(path: Path, cfg: ChunkConfig) -> ParsedDoc:
    raw = path.read_text(encoding="utf-8", errors="replace")
    blocks: list[Block] = []
    title = path.stem
    for chunk in re.split(r"\n\s*\n", raw):
        lines = [ln for ln in chunk.splitlines() if ln.strip()]
        if not lines:
            continue
        if all(ln.lstrip().startswith("|") for ln in lines):  # markdown table
            rows = [
                [c.strip() for c in ln.strip().strip("|").split("|")]
                for ln in lines
                if not re.fullmatch(r"[\s|:-]+", ln)
            ]
            rows = _clean_rows(rows)
            if len(rows) >= 2:
                header, body = split_table_header(rows)
                blocks.append(Block("table", header=header, rows=body))
                continue
        first = clean_text(lines[0])
        md = re.match(r"^(#{1,6})\s+(.*)", first)
        if md and len(lines) == 1:
            blocks.append(Block("heading", md.group(2).strip(), level=len(md.group(1))))
        elif len(lines) == 1 and looks_like_heading(first, plain=True):
            blocks.append(Block("heading", first, level=numbering_depth(first) or 1))
        elif all(_LIST_MARK.match(ln) for ln in lines):
            blocks += [Block("_li", clean_text(ln)) for ln in lines]
        else:
            blocks.append(Block("paragraph", join_lines(lines)))
    if blocks and blocks[0].kind == "heading" and not _HEAD_NUM.match(blocks[0].text):
        first = blocks[0]
        if (
            first.text.isupper()
            or sum(b.kind == "heading" and b.level <= first.level for b in blocks) == 1
        ):
            title, blocks = first.text, blocks[1:]
    return ParsedDoc(
        title, _attach_captions(_group_lists(blocks)), None, {"extractor": "text"}
    )


def extract(path: Path, cfg: ChunkConfig) -> ParsedDoc:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_pdf(path, cfg)
    if ext == ".docx":
        return extract_docx(path, cfg)
    if ext in (".txt", ".md", ".markdown"):
        return extract_text(path, cfg)
    raise ValueError(f"Unsupported file type: {ext}")
